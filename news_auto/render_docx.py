from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from .models import NewsItem
from .official_source import detect_official_source


def render_word_draft(
    items: Iterable[NewsItem],
    template_path: Path,
    output_path: Path,
    authority_tokens: Iterable[str] = (),
    non_official_warning: str = "【可能非官方网站，待核实修改】",
) -> Path:
    if template_path.exists():
        document = Document(str(template_path))
        clear_body(document)
    else:
        document = Document()
    for index, item in enumerate(items, start=1):
        title = item.title or "未获取标题"
        add_paragraph(document, f"{index}. {title}", bold=True)
        add_paragraph(document, f"摘要：{item.summary or '[待生成摘要]'}")
        link = item.official_link or item.link or "[未获取链接]"
        is_official, _ = detect_official_source(link, authority_tokens)
        warning = "" if is_official else non_official_warning
        add_paragraph(document, f"原文链接：{link}{warning}")
        add_paragraph(document, "")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def clear_body(document: Document) -> None:
    body = document._body._element  # noqa: SLF001 - python-docx has no public clear API
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def add_paragraph(document: Document, text: str, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    apply_run_font(run)


def apply_run_font(run) -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")  # noqa: SLF001
    run.font.size = Pt(11)