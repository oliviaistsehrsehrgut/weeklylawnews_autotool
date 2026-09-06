from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.oxml.ns import qn

from .models import NewsItem

TITLE_PREFIX = re.compile(r"^\s*\d+\s*[..．、]\s*")
SUMMARY_PREFIX = re.compile(r"^\s*摘要\s*[::：]\s*")
LINK_PREFIX = re.compile(r"^\s*原文链接\s*[::：]\s*")


def parse_reviewed_docx(path: Path) -> List[NewsItem]:
    return parse_document(Document(str(path)))


def parse_reviewed_docx_bytes(data: bytes) -> List[NewsItem]:
    return parse_document(Document(BytesIO(data)))


def _has_auto_numbering(paragraph) -> bool:
    """Return True if Word auto-numbers this paragraph (List Paragraph with w:numPr)."""
    pPr = paragraph._element.find(qn("w:pPr"))
    if pPr is None:
        return False
    return pPr.find(qn("w:numPr")) is not None


def _is_title(paragraph, text: str) -> bool:
    """True if this paragraph is a news item title."""
    # Explicitly typed number prefix: "1. ", "2、", etc.
    if TITLE_PREFIX.match(text):
        return True
    # Word auto-numbering (e.g. 法讯草稿模板.docx List Paragraph slots)
    return _has_auto_numbering(paragraph)


def parse_document(document: Document) -> List[NewsItem]:
    items: List[NewsItem] = []
    current: Optional[NewsItem] = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if SUMMARY_PREFIX.match(text):
            current = ensure_current(current)
            current.summary = SUMMARY_PREFIX.sub("", text).strip()
            continue
        if LINK_PREFIX.match(text):
            current = ensure_current(current)
            current.link = clean_word_link(LINK_PREFIX.sub("", text).strip())
            continue

        if _is_title(paragraph, text):
            if current and (current.title or current.summary or current.link or current.content):
                items.append(current)
            current = NewsItem(title=TITLE_PREFIX.sub("", text).strip(), source="人工审查 Word")
        elif current is None:
            current = NewsItem(title=text, source="人工审查 Word")
        else:
            current.content = "\n".join(part for part in (current.content, text) if part)

    if current and (current.title or current.summary or current.link or current.content):
        items.append(current)
    return [item for item in items if item.title or item.summary or item.link or item.content]


def clean_word_link(link: str) -> str:
    return link.replace("​", "")


def ensure_current(current: Optional[NewsItem]) -> NewsItem:
    if current is None:
        return NewsItem(source="人工审查 Word")
    return current
